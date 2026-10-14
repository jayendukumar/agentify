"""A small deck kit: slides are plain data, rendered to a self-contained HTML
deck (keyboard/click navigation, same look as the Meridian case study) and to
PowerPoint from the same data, so the two never drift.

    from deck_kit import *
    slides = [title("Kicker", "Title", "Subtitle"), cards(...), ...]
    build(slides, "deck.html", "deck.pptx", title="...")

Slide kinds: title, section, cards, compare, table, kpis, bars, roadmap, flow,
shot, statement. Colours: the existing deck navy/blue; series colours for
as-is vs to-be are the dataviz-validated pair (#c2601d, #006ffb).
"""

from __future__ import annotations

import base64
import html
import io
from pathlib import Path
from typing import Any

E = html.escape
AS_IS, TO_BE = "#c2601d", "#006ffb"

_TONE_LABEL = {"bad": "Critical", "warn": "High", "good": "Strength"}

# ---- slide constructors -------------------------------------------------------


def title(kicker: str, heading: str, sub: str, hint: str = "") -> dict:
    return {"kind": "title", "kicker": kicker, "title": heading, "sub": sub, "hint": hint}


def section(kicker: str, heading: str, sub: str = "") -> dict:
    return {"kind": "section", "kicker": kicker, "title": heading, "sub": sub}


def cards(kicker: str, heading: str, intro: str, items: list[tuple[str, str]], note: str = "") -> dict:
    return {"kind": "cards", "kicker": kicker, "title": heading, "intro": intro, "cards": items, "note": note}


def compare(kicker: str, heading: str, left_h: str, left: list[str], right_h: str, right: list[str], note: str = "") -> dict:
    return {"kind": "compare", "kicker": kicker, "title": heading, "left": (left_h, left), "right": (right_h, right), "note": note}


def table(kicker: str, heading: str, intro: str, headers: list[str], rows: list[list[str]], note: str = "",
          tone: list[str] | None = None, tone_labels: dict[str, str] | None = None) -> dict:
    """`tone`: per row "", "good", "warn" or "bad" (shown with a label, never colour alone);
    `tone_labels` renames the labels (default Critical / High / Strength)."""
    return {"kind": "table", "kicker": kicker, "title": heading, "intro": intro, "headers": headers, "rows": rows,
            "note": note, "tone": tone or [""] * len(rows), "tone_labels": {**_TONE_LABEL, **(tone_labels or {})}}


def kpis(kicker: str, heading: str, intro: str, tiles: list[tuple[str, str, str]], note: str = "") -> dict:
    return {"kind": "kpis", "kicker": kicker, "title": heading, "intro": intro, "tiles": tiles, "note": note}


def bars(kicker: str, heading: str, intro: str, rows: list[tuple[str, float, float, str, str]], note: str = "",
         legend: tuple[str, str] = ("Today", "With agents")) -> dict:
    """rows: (measure, as-is value, to-be value, as-is label, to-be label). Each pair shares one scale."""
    return {"kind": "bars", "kicker": kicker, "title": heading, "intro": intro, "rows": rows, "note": note, "legend": legend}


def roadmap(kicker: str, heading: str, intro: str, phases: list[tuple[str, str, list[str]]], note: str = "") -> dict:
    return {"kind": "roadmap", "kicker": kicker, "title": heading, "intro": intro, "phases": phases, "note": note}


def flow(kicker: str, heading: str, intro: str, steps: list[tuple[str, str, str]], note: str = "") -> dict:
    """steps: (label, sub-label, kind) with kind in event/agent/person/gate."""
    return {"kind": "flow", "kicker": kicker, "title": heading, "intro": intro, "steps": steps, "note": note}


def shot(kicker: str, heading: str, image: Path, bullets: list[str], why: str = "") -> dict:
    return {"kind": "shot", "kicker": kicker, "title": heading, "image": Path(image), "bullets": bullets, "why": why}


def statement(kicker: str, heading: str, body: str, points: list[str] | None = None) -> dict:
    return {"kind": "statement", "kicker": kicker, "title": heading, "body": body, "points": points or []}


# ---- HTML -----------------------------------------------------------------------

CSS = """
*{box-sizing:border-box} html,body{margin:0;height:100%;background:#0d1530;font-family:Inter,Segoe UI,Arial,sans-serif}
.slide{display:none;position:absolute;inset:0;margin:auto;width:100vw;height:56.25vw;max-height:100vh;max-width:177.78vh;overflow:hidden}
.slide.on{display:block} .slide.shot-slide.on{display:grid;grid-template-columns:68% 32%}
.num{position:absolute;bottom:2.4vmin;right:3vmin;font-size:1.6vmin;color:#8fb3ff} .light .num{color:#5d6a90}
.kicker{color:#8fb3ff;font-size:1.8vmin;font-weight:700;letter-spacing:.06em;text-transform:uppercase} .kicker.dark{color:#2f5bd3}
.hero{background:linear-gradient(135deg,#16245c,#0d1530);color:#fff}
.hero .inner{position:absolute;left:9%;right:9%;top:50%;transform:translateY(-50%)}
.hero h1{font-size:6.2vmin;line-height:1.1;margin:2vmin 0} .hero .sub{font-size:2.7vmin;color:#dfe7ff;margin:0;line-height:1.45}
.hero .hint{font-size:1.8vmin;color:#8fb3ff;margin-top:6vmin} .hero.section h1{font-size:7.4vmin}
.light{background:#f6f8fc;color:#16245c} .light .inner{position:absolute;inset:4.5vmin 6.5vmin 6.5vmin;display:flex;flex-direction:column;justify-content:center}
.light h2{font-size:4.6vmin;margin:1vmin 0 1.6vmin;line-height:1.15} .intro{font-size:2.5vmin;line-height:1.45;color:#3b4a75;max-width:90%;margin:0 0 2.4vmin}
.note{margin-top:2vmin;font-size:1.75vmin;color:#5d6a90;font-style:italic}
.cards{display:grid;gap:2.4vmin} .cards.n2,.cards.n4{grid-template-columns:1fr 1fr} .cards.n3,.cards.n6,.cards.n5{grid-template-columns:1fr 1fr 1fr}
.card{background:#fff;border-radius:1.4vmin;padding:2.8vmin;box-shadow:0 .4vmin 1.6vmin rgba(22,36,92,.08);border-top:.6vmin solid #2f5bd3}
.card h3{margin:0 0 1.1vmin;font-size:2.8vmin} .card p{margin:0;font-size:2.2vmin;line-height:1.45;color:#3b4a75}
.versus{display:grid;grid-template-columns:1fr 1fr;gap:3vmin}
.col{border-radius:1.4vmin;padding:3.2vmin} .col h3{margin:0 0 1.6vmin;font-size:3vmin}
.col ul{margin:0;padding-left:2.6vmin} .col li{font-size:2.35vmin;line-height:1.4;margin-bottom:1.4vmin}
.before{background:#eceef4;color:#4a5370} .after{background:#16245c;color:#fff} .after li{color:#dfe7ff}
table.t{border-collapse:collapse;width:100%;background:#fff;border-radius:1.2vmin;overflow:hidden;box-shadow:0 .4vmin 1.6vmin rgba(22,36,92,.08)}
table.t th{background:#16245c;color:#fff;text-align:left;font-size:1.9vmin;padding:1.2vmin 1.4vmin;font-weight:600}
table.t td{font-size:1.95vmin;padding:1.05vmin 1.4vmin;border-top:1px solid #e3e8f4;color:#2b3760;vertical-align:top;line-height:1.35}
table.t tr.bad td:first-child{border-left:.6vmin solid #b42318} table.t tr.warn td:first-child{border-left:.6vmin solid #b54708}
table.t tr.good td:first-child{border-left:.6vmin solid #067647}
.tag{display:inline-block;font-size:1.5vmin;font-weight:700;border-radius:.6vmin;padding:.2vmin .8vmin;margin-left:.8vmin}
.tag.bad{background:#fef3f2;color:#b42318} .tag.warn{background:#fffaeb;color:#b54708} .tag.good{background:#ecfdf3;color:#067647}
.tiles{display:grid;grid-template-columns:repeat(auto-fit,minmax(22vmin,1fr));gap:2.2vmin}
.tile{background:#fff;border-radius:1.4vmin;padding:2.6vmin;box-shadow:0 .4vmin 1.6vmin rgba(22,36,92,.08)}
.tile .l{font-size:1.8vmin;color:#5d6a90;text-transform:uppercase;letter-spacing:.04em;font-weight:700}
.tile .v{font-size:5.4vmin;font-weight:800;color:#16245c;margin:.8vmin 0;font-variant-numeric:tabular-nums}
.tile .n{font-size:1.8vmin;color:#3b4a75;line-height:1.35}
.legend{display:flex;gap:3vmin;font-size:2vmin;color:#3b4a75;margin-bottom:1.6vmin} .sw{display:inline-block;width:1.8vmin;height:1.8vmin;border-radius:.4vmin;vertical-align:middle;margin-right:.8vmin}
.pair{background:#fff;border-radius:1.2vmin;padding:2vmin 2.4vmin;margin-bottom:1.6vmin;box-shadow:0 .4vmin 1.6vmin rgba(22,36,92,.06)}
.pair h4{margin:0 0 1vmin;font-size:2.2vmin;color:#16245c}
.prow{display:grid;grid-template-columns:16vmin 1fr 18vmin;align-items:center;gap:1.6vmin;margin:.7vmin 0;font-size:1.9vmin;color:#3b4a75}
.ptrack{height:2.2vmin} .pbar{height:2.2vmin;border-radius:0 .5vmin .5vmin 0} .pval{text-align:right;font-variant-numeric:tabular-nums;font-weight:600}
.phases{display:grid;gap:2vmin} .phase{background:#fff;border-radius:1.2vmin;padding:2.2vmin 2.4vmin;box-shadow:0 .4vmin 1.6vmin rgba(22,36,92,.08);border-top:.6vmin solid #2f5bd3}
.phase h3{margin:0;font-size:2.5vmin} .phase .d{font-size:1.7vmin;color:#2f5bd3;font-weight:700;margin:.4vmin 0 1vmin} .phase ul{margin:0;padding-left:2.2vmin} .phase li{font-size:1.85vmin;line-height:1.38;margin-bottom:.6vmin;color:#3b4a75}
.fl{display:flex;flex-wrap:wrap;align-items:center;gap:1vmin 0}
.fs{border-radius:1vmin;padding:1.2vmin 1.4vmin;min-width:15vmin;max-width:20vmin;text-align:center;font-size:1.75vmin;line-height:1.3;border:.3vmin solid #c9d3ee;background:#fff}
.fs b{display:block;font-size:1.85vmin;color:#16245c} .fs span{font-size:1.5vmin;color:#5d6a90}
.fs.agent{border-color:#006ffb;background:#edf4ff} .fs.person{border-color:#c2601d;background:#fff4ec} .fs.event{border-radius:5vmin;background:#eef2fb}
.fs.gate{border-style:dashed;border-color:#5d6a90;background:#f6f8fc}
.fa{color:#2f5bd3;font-size:2.4vmin;padding:0 .6vmin}
.shot{display:flex;align-items:center;justify-content:center;padding:1.6vmin;background:#e9edf6}
.shot img{max-width:100%;max-height:100%;border-radius:.8vmin;box-shadow:0 .5vmin 2vmin rgba(13,21,48,.25)}
aside{position:relative;background:#16245c;color:#fff;padding:4.2vmin 3vmin}
aside h2{font-size:3.2vmin;line-height:1.2;margin:1.4vmin 0 2vmin} aside ul{padding-left:2.4vmin;margin:0}
aside li{font-size:2vmin;line-height:1.42;margin-bottom:1.3vmin;color:#dfe7ff}
.why{margin-top:2.2vmin;background:rgba(143,179,255,.14);border-left:.5vmin solid #8fb3ff;border-radius:.8vmin;padding:1.5vmin 1.7vmin;font-size:1.9vmin;line-height:1.42}
.why b{display:block;color:#8fb3ff;font-size:1.5vmin;text-transform:uppercase;letter-spacing:.06em;margin-bottom:.5vmin}
.stmt .body{font-size:3vmin;line-height:1.5;color:#2b3760;max-width:88%} .stmt ul{margin-top:2.4vmin;padding-left:3vmin} .stmt li{font-size:2.4vmin;line-height:1.45;margin-bottom:1.2vmin;color:#3b4a75}
"""

JS = """
const s=[...document.querySelectorAll('.slide')]; let i=0;
const show=n=>{i=Math.max(0,Math.min(s.length-1,n)); s.forEach((e,k)=>e.classList.toggle('on',k===i)); location.hash=i+1};
document.addEventListener('keydown',e=>{if(['ArrowRight','PageDown',' '].includes(e.key))show(i+1);
  if(['ArrowLeft','PageUp'].includes(e.key))show(i-1); if(e.key==='Home')show(0); if(e.key==='End')show(s.length-1);
  if(e.key==='f'||e.key==='F')document.documentElement.requestFullscreen?.()});
document.addEventListener('click',e=>{if(e.target.closest('a'))return; show(e.clientX<innerWidth/4?i-1:i+1)});
const fromHash=()=>{const n=(parseInt(location.hash.slice(1))||1)-1; if(n!==i||!s[i].classList.contains('on'))show(n)};
addEventListener('hashchange',fromHash); fromHash();
"""


def _img_uri(path: Path) -> str:
    from PIL import Image

    im = Image.open(path).convert("RGB")
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=86)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


def _light(s: dict, body: str) -> str:
    intro = f'<p class="intro">{E(s["intro"])}</p>' if s.get("intro") else ""
    note = f'<p class="note">{E(s["note"])}</p>' if s.get("note") else ""
    return (f'<section class="slide light"><div class="inner"><div class="kicker dark">{E(s["kicker"])}</div>'
            f'<h2>{E(s["title"])}</h2>{intro}{body}{note}</div></section>')


def _html_slide(s: dict) -> str:
    k = s["kind"]
    if k in ("title", "section"):
        hint = f'<p class="hint">{E(s["hint"])}</p>' if s.get("hint") else ""
        return (f'<section class="slide hero{" section" if k == "section" else ""}"><div class="inner">'
                f'<div class="kicker">{E(s["kicker"])}</div><h1>{E(s["title"])}</h1><p class="sub">{E(s["sub"])}</p>{hint}</div></section>')
    if k == "cards":
        cs = "".join(f'<div class="card"><h3>{E(h)}</h3><p>{E(t)}</p></div>' for h, t in s["cards"])
        return _light(s, f'<div class="cards n{len(s["cards"])}">{cs}</div>')
    if k == "compare":
        li = lambda xs: "".join(f"<li>{E(x)}</li>" for x in xs)  # noqa: E731
        (lh, lx), (rh, rx) = s["left"], s["right"]
        return _light(s, f'<div class="versus"><div class="col before"><h3>{E(lh)}</h3><ul>{li(lx)}</ul></div>'
                         f'<div class="col after"><h3>{E(rh)}</h3><ul>{li(rx)}</ul></div></div>')
    if k == "table":
        head = "".join(f"<th>{E(h)}</th>" for h in s["headers"])
        body = ""
        for row, tone in zip(s["rows"], s["tone"]):
            cells = [E(str(c)) for c in row]
            if tone:
                cells[0] += f'<span class="tag {tone}">{E(s["tone_labels"].get(tone, tone))}</span>'
            body += f'<tr class="{tone}">' + "".join(f"<td>{c}</td>" for c in cells) + "</tr>"
        return _light(s, f'<table class="t"><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>')
    if k == "kpis":
        ts = "".join(f'<div class="tile"><div class="l">{E(lab)}</div><div class="v">{E(v)}</div><div class="n">{E(n)}</div></div>'
                     for lab, v, n in s["tiles"])
        return _light(s, f'<div class="tiles">{ts}</div>')
    if k == "bars":
        a, b = s["legend"]
        out = (f'<div class="legend"><span><span class="sw" style="background:{AS_IS}"></span>{E(a)}</span>'
               f'<span><span class="sw" style="background:{TO_BE}"></span>{E(b)}</span></div>')
        for measure, va, vb, la, lb in s["rows"]:
            m = max(va, vb, 1e-9)
            out += f'<div class="pair"><h4>{E(measure)}</h4>'
            for lab, v, txt, col in ((a, va, la, AS_IS), (b, vb, lb, TO_BE)):
                out += (f'<div class="prow" title="{E(lab)}: {E(txt)}"><span>{E(lab)}</span><span class="ptrack">'
                        f'<div class="pbar" style="width:{max(0.8, v / m * 100):.1f}%;background:{col}"></div></span>'
                        f'<span class="pval">{E(txt)}</span></div>')
            out += "</div>"
        return _light(s, out)
    if k == "roadmap":
        ph = "".join(f'<div class="phase"><h3>{E(n)}</h3><div class="d">{E(d)}</div><ul>{"".join(f"<li>{E(x)}</li>" for x in xs)}</ul></div>'
                     for n, d, xs in s["phases"])
        return _light(s, f'<div class="phases" style="grid-template-columns:repeat({len(s["phases"])},1fr)">{ph}</div>')
    if k == "flow":
        items = []
        for i, (lab, sub, kind) in enumerate(s["steps"]):
            items.append(f'<div class="fs {kind}"><b>{E(lab)}</b>{f"<span>{E(sub)}</span>" if sub else ""}</div>')
            if i < len(s["steps"]) - 1:
                items.append('<div class="fa">&#10140;</div>')
        legend = (f'<div class="legend" style="margin-top:2vmin"><span><span class="sw" style="background:#edf4ff;border:2px solid #006ffb"></span>Agent</span>'
                  f'<span><span class="sw" style="background:#fff4ec;border:2px solid #c2601d"></span>Person</span>'
                  f'<span><span class="sw" style="background:#f6f8fc;border:2px dashed #5d6a90"></span>Decision point</span></div>')
        return _light(s, f'<div class="fl">{"".join(items)}</div>{legend}')
    if k == "shot":
        lis = "".join(f"<li>{E(x)}</li>" for x in s["bullets"])
        why = f'<div class="why"><b>Why it matters</b>{E(s["why"])}</div>' if s.get("why") else ""
        return (f'<section class="slide shot-slide"><div class="shot"><img src="{_img_uri(s["image"])}" alt="{E(s["title"])}"></div>'
                f'<aside><div class="kicker">{E(s["kicker"])}</div><h2>{E(s["title"])}</h2><ul>{lis}</ul>{why}</aside></section>')
    if k == "statement":
        pts = "".join(f"<li>{E(p)}</li>" for p in s["points"])
        return (f'<section class="slide light stmt"><div class="inner"><div class="kicker dark">{E(s["kicker"])}</div>'
                f'<h2>{E(s["title"])}</h2><p class="body">{E(s["body"])}</p>{f"<ul>{pts}</ul>" if pts else ""}</div></section>')
    raise ValueError(k)


def build_html(slides: list[dict], out: Path, doc_title: str) -> None:
    rendered = [_html_slide(s) for s in slides]
    numbered = [r.replace("</section>", f'<div class="num">{i + 1} / {len(rendered)}</div></section>', 1)
                if i else r for i, r in enumerate(rendered)]
    out.write_text(f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{E(doc_title)}</title>
<meta name="viewport" content="width=device-width,initial-scale=1"><style>{CSS}</style></head><body>
{chr(10).join(numbered)}
<script>{JS}</script></body></html>""", encoding="utf-8")


# ---- PowerPoint -----------------------------------------------------------------


def build_pptx(slides: list[dict], out: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    rgb = lambda h: RGBColor.from_string(h.lstrip("#"))  # noqa: E731
    NAVY, DEEP, BLUE, SKY, LIGHT, INK, MUTED, WHITE = (rgb(c) for c in (
        "16245c", "0d1530", "2f5bd3", "8fb3ff", "f6f8fc", "3b4a75", "5d6a90", "ffffff"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def bg(slide, color):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def text(slide, x, y, w, h, value, size, color=INK, bold=False, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        lines = value if isinstance(value, list) else [value]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = line
            run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(size), bold, color, "Segoe UI"
        return box

    def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        r = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        r.fill.solid()
        r.fill.fore_color.rgb = fill
        if line:
            r.line.color.rgb = line
        else:
            r.line.fill.background()
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            r.adjustments[0] = 0.08
        return r

    def header(slide, s, top=0.45):
        bg(slide, LIGHT)
        text(slide, 0.7, top, 12, 0.35, s["kicker"].upper(), 12, BLUE, True)
        text(slide, 0.7, top + 0.32, 12, 0.8, s["title"], 28, NAVY, True)
        y = top + 1.15
        if s.get("intro"):
            text(slide, 0.7, y, 11.8, 0.8, s["intro"], 15, INK)
            y += 0.85
        if s.get("note"):
            text(slide, 0.7, 6.85, 11.8, 0.4, s["note"], 11, MUTED)
        return y

    for s in slides:
        slide = prs.slides.add_slide(blank)
        k = s["kind"]
        if k in ("title", "section"):
            bg(slide, DEEP if k == "title" else NAVY)
            text(slide, 1.2, 2.2, 11, 0.4, s["kicker"].upper(), 14, SKY, True)
            text(slide, 1.2, 2.7, 11, 1.6, s["title"], 40 if k == "title" else 44, WHITE, True)
            text(slide, 1.2, 4.4, 11, 1.4, s["sub"], 18, rgb("dfe7ff"))
            continue
        if k == "shot":
            bg(slide, rgb("e9edf6"))
            from PIL import Image

            im = Image.open(s["image"])
            iw, ih = im.size
            box_w, box_h = 8.6, 7.0
            scale = min(box_w / iw, box_h / ih)
            slide.shapes.add_picture(str(s["image"]), Inches(0.25 + (box_w - iw * scale) / 2), Inches(0.25 + (box_h - ih * scale) / 2),
                                     Inches(iw * scale), Inches(ih * scale))
            rect(slide, 9.1, 0, 4.233, 7.5, NAVY, shape=MSO_SHAPE.RECTANGLE)
            text(slide, 9.35, 0.5, 3.8, 0.3, s["kicker"].upper(), 11, SKY, True)
            text(slide, 9.35, 0.85, 3.8, 1.2, s["title"], 20, WHITE, True)
            text(slide, 9.35, 2.2, 3.8, 3.4, [f"• {b}" for b in s["bullets"]], 12, rgb("dfe7ff"))
            if s.get("why"):
                text(slide, 9.35, 5.7, 3.8, 1.5, ["WHY IT MATTERS", s["why"]], 11, SKY)
            continue
        y = header(slide, s)
        if k == "cards":
            n = len(s["cards"])
            cols = 2 if n in (2, 4) else 3
            rows_n = (n + cols - 1) // cols
            cw, ch = (11.9 - 0.3 * (cols - 1)) / cols, min(2.3, (6.7 - y) / rows_n - 0.2)
            for i, (h, t) in enumerate(s["cards"]):
                x0, y0 = 0.7 + (i % cols) * (cw + 0.3), y + (i // cols) * (ch + 0.25)
                rect(slide, x0, y0, cw, ch, WHITE)
                rect(slide, x0, y0, cw, 0.07, BLUE, shape=MSO_SHAPE.RECTANGLE)
                text(slide, x0 + 0.2, y0 + 0.15, cw - 0.4, 0.5, h, 16, NAVY, True)
                text(slide, x0 + 0.2, y0 + 0.62, cw - 0.4, ch - 0.7, t, 12, INK)
        elif k == "compare":
            for i, ((h, xs), fill, col) in enumerate(((s["left"], rgb("eceef4"), INK), (s["right"], NAVY, rgb("dfe7ff")))):
                x0 = 0.7 + i * 6.1
                rect(slide, x0, y, 5.8, 6.6 - y, fill)
                text(slide, x0 + 0.3, y + 0.2, 5.2, 0.5, h, 18, WHITE if i else NAVY, True)
                text(slide, x0 + 0.3, y + 0.8, 5.2, 5.6 - y, [f"• {x}" for x in xs], 13, col)
        elif k == "table":
            rows = s["rows"]
            shape = slide.shapes.add_table(len(rows) + 1, len(s["headers"]), Inches(0.7), Inches(y), Inches(11.9),
                                           Inches(min(6.7 - y, 0.42 * (len(rows) + 1))))
            tbl = shape.table
            for c, h in enumerate(s["headers"]):
                cell = tbl.cell(0, c)
                cell.text = h
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size, r.font.bold, r.font.color.rgb = Pt(11), True, WHITE
            for r_i, (row, tone) in enumerate(zip(rows, s["tone"]), 1):
                for c, v in enumerate(row):
                    cell = tbl.cell(r_i, c)
                    cell.text = str(v) + (f"  [{s['tone_labels'][tone]}]" if c == 0 and tone in s["tone_labels"] else "")
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = WHITE
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size, r.font.color.rgb = Pt(10), INK
        elif k == "kpis":
            n = len(s["tiles"])
            cols = min(n, 4)
            tw = (11.9 - 0.25 * (cols - 1)) / cols
            for i, (lab, v, note) in enumerate(s["tiles"]):
                x0, y0 = 0.7 + (i % cols) * (tw + 0.25), y + (i // cols) * 2.2
                rect(slide, x0, y0, tw, 2.0, WHITE)
                text(slide, x0 + 0.2, y0 + 0.15, tw - 0.4, 0.3, lab.upper(), 10, MUTED, True)
                text(slide, x0 + 0.2, y0 + 0.45, tw - 0.4, 0.8, v, 30, NAVY, True)
                text(slide, x0 + 0.2, y0 + 1.25, tw - 0.4, 0.7, note, 10, INK)
        elif k == "bars":
            a, b = s["legend"]
            rect(slide, 0.7, y, 0.2, 0.2, rgb(AS_IS), shape=MSO_SHAPE.RECTANGLE)
            text(slide, 0.95, y - 0.07, 2, 0.3, a, 12, INK)
            rect(slide, 2.6, y, 0.2, 0.2, rgb(TO_BE), shape=MSO_SHAPE.RECTANGLE)
            text(slide, 2.85, y - 0.07, 3, 0.3, b, 12, INK)
            y += 0.45
            for measure, va, vb, la, lb in s["rows"]:
                m = max(va, vb, 1e-9)
                text(slide, 0.7, y, 11, 0.3, measure, 13, NAVY, True)
                for j, (lab, v, txt, col) in enumerate(((a, va, la, AS_IS), (b, vb, lb, TO_BE))):
                    yy = y + 0.38 + j * 0.34
                    text(slide, 0.7, yy - 0.05, 1.8, 0.3, lab, 11, INK)
                    rect(slide, 2.6, yy, max(0.05, 8.0 * v / m), 0.24, rgb(col), shape=MSO_SHAPE.RECTANGLE)
                    text(slide, 10.8, yy - 0.05, 1.8, 0.3, txt, 11, INK, True)
                y += 1.15
        elif k == "roadmap":
            n = len(s["phases"])
            pw = (11.9 - 0.2 * (n - 1)) / n
            for i, (name, dur, xs) in enumerate(s["phases"]):
                x0 = 0.7 + i * (pw + 0.2)
                rect(slide, x0, y, pw, 6.7 - y, WHITE)
                rect(slide, x0, y, pw, 0.07, BLUE, shape=MSO_SHAPE.RECTANGLE)
                text(slide, x0 + 0.15, y + 0.12, pw - 0.3, 0.5, name, 14, NAVY, True)
                text(slide, x0 + 0.15, y + 0.62, pw - 0.3, 0.3, dur, 10, BLUE, True)
                text(slide, x0 + 0.15, y + 0.95, pw - 0.3, 5.5 - y, [f"• {x}" for x in xs], 10, INK)
        elif k == "flow":
            colours = {"agent": ("edf4ff", "006ffb"), "person": ("fff4ec", "c2601d"), "event": ("eef2fb", "c9d3ee"), "gate": ("f6f8fc", "5d6a90")}
            per_row = 6
            bw = 1.8
            for i, (lab, sub, kind) in enumerate(s["steps"]):
                x0, y0 = 0.7 + (i % per_row) * 2.05, y + (i // per_row) * 1.35
                fill, line = colours.get(kind, colours["event"])
                rect(slide, x0, y0, bw, 1.05, rgb(fill), rgb(line))
                text(slide, x0 + 0.08, y0 + 0.05, bw - 0.16, 0.95, [lab] + ([sub] if sub else []), 9, NAVY, True, MSO_ANCHOR.MIDDLE)
        elif k == "statement":
            text(slide, 0.7, y, 11.8, 2.0, s["body"], 18, INK)
            if s["points"]:
                text(slide, 0.7, y + 2.1, 11.8, 6.6 - y - 2.1, [f"• {p}" for p in s["points"]], 14, INK)
    prs.save(out)


def build(slides: list[dict], html_out: Path, pptx_out: Path | None, *, title: str) -> None:
    build_html(slides, Path(html_out), title)
    if pptx_out:
        build_pptx(slides, Path(pptx_out))
    print(f"{html_out}: {len(slides)} slides" + (f"; {pptx_out}" if pptx_out else ""))
