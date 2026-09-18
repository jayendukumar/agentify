import io
from unittest.mock import patch

from docx import Document as DocxDocument

from app.ingestion.extractors import extract_docx, extract_pdf


def _build_docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Step 1: Submit the request form.")
    document.add_paragraph("")  # blank paragraph -- should be skipped
    document.add_paragraph("Step 2: Manager reviews the request.")

    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Owner"
    table.rows[1].cells[1].text = "Finance"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_docx_returns_text_and_table_blocks():
    blocks = extract_docx(_build_docx_bytes())

    text_blocks = [b for b in blocks if b.kind == "text"]
    table_blocks = [b for b in blocks if b.kind == "table"]

    assert [b.content for b in text_blocks] == [
        "Step 1: Submit the request form.",
        "Step 2: Manager reviews the request.",
    ]
    assert text_blocks[0].location == "paragraph 1"
    assert text_blocks[1].location == "paragraph 3"

    assert len(table_blocks) == 1
    assert table_blocks[0].location == "table 1"
    assert "Field | Value" in table_blocks[0].content
    assert "Owner | Finance" in table_blocks[0].content


def test_extract_docx_skips_empty_document():
    blocks = extract_docx(_build_docx_bytes_empty())
    assert blocks == []


def _build_docx_bytes_empty() -> bytes:
    document = DocxDocument()
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class _FakePage:
    def __init__(self, text: str, tables: list[list[list[str | None]]]):
        self._text = text
        self._tables = tables

    def extract_text(self) -> str:
        return self._text

    def extract_tables(self) -> list[list[list[str | None]]]:
        return self._tables


class _FakePdf:
    def __init__(self, pages: list[_FakePage]):
        self.pages = pages

    def __enter__(self) -> "_FakePdf":
        return self

    def __exit__(self, *args: object) -> None:
        return None


def test_extract_pdf_returns_text_and_table_blocks_per_page():
    fake_pdf = _FakePdf(
        pages=[
            _FakePage("Step 1: Submit the request form.", tables=[]),
            _FakePage("", tables=[[["Field", "Value"], ["Owner", "Finance"]]]),
        ]
    )

    with patch("app.ingestion.extractors.pdfplumber.open", return_value=fake_pdf):
        blocks = extract_pdf(b"fake-pdf-bytes")

    text_blocks = [b for b in blocks if b.kind == "text"]
    table_blocks = [b for b in blocks if b.kind == "table"]

    assert len(text_blocks) == 1
    assert text_blocks[0].location == "page 1"
    assert text_blocks[0].content == "Step 1: Submit the request form."

    assert len(table_blocks) == 1
    assert table_blocks[0].location == "page 2, table 1"
    assert "Field | Value" in table_blocks[0].content
