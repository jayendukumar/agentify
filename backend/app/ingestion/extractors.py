"""Deterministic extraction for US1.3 (PDF/DOCX) and US1.4 (Visio).

No chunking: callers get back every block from the whole document. See
planning/document-ingestion-strategy.md for why chunking is an explicit
fast-follow rather than built here.
"""

from __future__ import annotations

import io
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from typing import Literal

import pdfplumber
from docx import Document as DocxDocument

BlockKind = Literal["text", "table"]


@dataclass
class ExtractedBlock:
    kind: BlockKind
    location: str
    content: str


def extract_pdf(data: bytes) -> list[ExtractedBlock]:
    blocks: list[ExtractedBlock] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_index, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                blocks.append(ExtractedBlock(kind="text", location=f"page {page_index}", content=text))

            for table_index, table in enumerate(page.extract_tables(), start=1):
                rendered = _render_table(table)
                if rendered:
                    blocks.append(
                        ExtractedBlock(
                            kind="table", location=f"page {page_index}, table {table_index}", content=rendered
                        )
                    )
    return blocks


def extract_docx(data: bytes) -> list[ExtractedBlock]:
    blocks: list[ExtractedBlock] = []
    document = DocxDocument(io.BytesIO(data))

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            blocks.append(ExtractedBlock(kind="text", location=f"paragraph {index}", content=text))

    for table_index, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        rendered = _render_table(rows)
        if rendered:
            blocks.append(ExtractedBlock(kind="table", location=f"table {table_index}", content=rendered))

    return blocks


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def extract_vsdx(data: bytes) -> list[ExtractedBlock]:
    """Parse a .vsdx (a zip of XML parts, per the process-doc-ingestion
    skill) deterministically: shape text becomes labeled blocks, and
    connectors become explicit "'X' -> 'Y'" relationship blocks derived
    from the file's actual topology rather than left for the LLM to guess
    from spatial layout.
    """
    blocks: list[ExtractedBlock] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        page_names = sorted(
            name
            for name in archive.namelist()
            if name.startswith("visio/pages/page") and name.endswith(".xml")
        )

        for page_index, page_name in enumerate(page_names, start=1):
            with archive.open(page_name) as f:
                root = ET.parse(f).getroot()

            shape_text: dict[str, str] = {}
            for shape in root.iter():
                if _local_name(shape.tag) != "Shape":
                    continue
                shape_id = shape.get("ID")
                if not shape_id:
                    continue
                text_el = next((c for c in shape if _local_name(c.tag) == "Text"), None)
                text = "".join(text_el.itertext()).strip() if text_el is not None else ""
                if text:
                    shape_text[shape_id] = text
                    blocks.append(
                        ExtractedBlock(kind="text", location=f"page {page_index}, shape {shape_id}", content=text)
                    )

            # A connector is its own shape; each of its two ends is a
            # separate <Connect> element sharing the connector's ID as
            # FromSheet, pointing at the two shapes it joins via ToSheet.
            connector_targets: dict[str, list[str]] = {}
            for connect in root.iter():
                if _local_name(connect.tag) != "Connect":
                    continue
                connector_id = connect.get("FromSheet")
                target_id = connect.get("ToSheet")
                if connector_id and target_id:
                    connector_targets.setdefault(connector_id, []).append(target_id)

            for connector_id, targets in connector_targets.items():
                unique_targets = list(dict.fromkeys(targets))
                if len(unique_targets) != 2:
                    continue
                from_id, to_id = unique_targets
                from_label = shape_text.get(from_id, f"shape {from_id}")
                to_label = shape_text.get(to_id, f"shape {to_id}")
                blocks.append(
                    ExtractedBlock(
                        kind="text",
                        location=f"page {page_index}, connector {connector_id}",
                        content=f"Connector: '{from_label}' -> '{to_label}'",
                    )
                )

    return blocks


def _render_table(rows: list[list[str | None]]) -> str:
    lines = []
    for row in rows:
        cells = [(cell or "").strip() for cell in row]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)
